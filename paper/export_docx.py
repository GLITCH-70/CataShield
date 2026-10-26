"""
export_docx.py
==============
Exports the IEEE paper (main.tex) to an editable Word document in IEEE conference layout
(US Letter, Times New Roman 10 pt, two columns, full-width figures/tables in one-column sections,
native Word equations).

Requires a compiled paper with intermediates (for citation, figure and table numbers):
    cd paper && ../tools/tectonic/tectonic.exe --keep-intermediates main.tex
then, from the project root:
    python paper/export_docx.py
Output: paper/main.docx
"""

from __future__ import annotations

import copy
import os
import re
import subprocess
import sys
import tempfile

import pypandoc
import pypdfium2 as pdfium
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TECTONIC = os.path.join(ROOT, "tools", "tectonic", "tectonic.exe")
WORK = tempfile.mkdtemp(prefix="docx_export_")
OUT = os.path.join(HERE, "main.docx")

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
WIDE_START, WIDE_END = "@@WIDESTART@@", "@@WIDEEND@@"
TITLE_MARK, AUTHOR_MARK, AFFIL_MARK, REF_MARK = "@@TITLE@@", "@@AUTHORS@@", "@@AFFIL@@", "@@REF@@"
FIGCAP_MARK, TABCAP_MARK, EQ_MARK = "@@FIGCAP@@", "@@TABCAP@@", "@@EQNUM"


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


# ----------------------------------------------------------------------------------------------
# 1. Assets: TikZ diagram and PDF figures -> PNG
# ----------------------------------------------------------------------------------------------
def render_pdf_png(pdf_path, png_path, scale=4.0):
    pdf = pdfium.PdfDocument(pdf_path)
    pdf[0].render(scale=scale).to_pil().save(png_path)


def build_tikz_png(tex):
    m = re.search(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", tex, re.S)
    if not m:
        return None
    standalone = ("\\documentclass[border=3pt]{standalone}\n\\usepackage{amsmath,amssymb}\n"
                  "\\usepackage{newtxtext,newtxmath}\n\\usepackage{tikz}\n\\usetikzlibrary{positioning,arrows.meta}\n"
                  "\\newcommand{\\eps}{\\varepsilon}\n\\begin{document}\n" + m.group(0) + "\n\\end{document}\n")
    src = os.path.join(WORK, "pipeline.tex")
    with open(src, "w", encoding="utf-8") as f:
        f.write(standalone)
    subprocess.run([TECTONIC, "pipeline.tex"], cwd=WORK, check=True, capture_output=True)
    png = os.path.join(WORK, "fig_pipeline.png")
    render_pdf_png(os.path.join(WORK, "pipeline.pdf"), png, scale=6.0)
    return png


# ----------------------------------------------------------------------------------------------
# 2. LaTeX preprocessing for pandoc
# ----------------------------------------------------------------------------------------------
def labels_and_cites():
    aux = read(os.path.join(HERE, "main.aux"))
    cites = dict(re.findall(r"\\bibcite\{([^}]+)\}\{(\d+)\}", aux))
    labels = {k: v for k, v in re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux)}
    return labels, {k: int(v) for k, v in cites.items()}


def format_cite(keys, cites):
    nums = sorted(cites[k.strip()] for k in keys.split(","))
    groups, start, prev = [], nums[0], nums[0]
    for n in nums[1:] + [None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        groups.append((start, prev))
        if n is not None:
            start = prev = n
    parts = []
    for a, b in groups:
        if b - a >= 2:
            parts.append(f"[{a}]--[{b}]")
        else:
            parts += [f"[{x}]" for x in range(a, b + 1)]
    return ", ".join(parts)


def bibliography_tex(cites):
    bbl = read(os.path.join(HERE, "main.bbl"))
    items = re.split(r"\\bibitem\{([^}]+)\}", bbl)[1:]
    entries = []
    for key, body in zip(items[0::2], items[1::2]):
        body = body.split("\\end{thebibliography}")[0]
        body = re.sub(r"\\BIBentry(ALT|STD)interwordspacing", "", body)
        body = re.sub(r"\\hskip[^\\]*\\relax", " ", body)
        body = body.replace("\\newblock", " ").replace("~", " ")
        body = re.sub(r"\s+", " ", body).strip()
        entries.append((cites[key], body))
    entries.sort()
    return "\n\n".join(f"{REF_MARK}[{n}] {b}" for n, b in entries)


def preprocess(tex, labels, cites, pipeline_png):
    # inline \input files (numbers + tables)
    tex = tex.replace("\\input{numbers}", read(os.path.join(HERE, "numbers.tex")))
    tex = re.sub(r"\\input\{(tables/[^}]+)\}", lambda m: read(os.path.join(HERE, m.group(1) + ".tex")), tex)
    # pandoc's tabular reader does not accept @{} column specs
    tex = tex.replace("@{}", "")
    # Word renders every $...$ as an equation object in Cambria Math. Plain numbers and the simple
    # fractions used in tables read better as ordinary text.
    tex = re.sub(r"\\frac\{1\}\{4\}|\\frac14", "¼", tex)
    tex = re.sub(r"\\frac\{1\}\{2\}|\\frac12", "½", tex)
    tex = re.sub(r"\\frac\{1\}\{255\}", "1/255", tex)
    tex = re.sub(r"\$([+\-]?\d+(?:\.\d+)?)\$", lambda m: m.group(1).replace("-", "\u2212"), tex)
    tex = re.sub(r"\$(¼|½)\$", r"\1", tex)
    tex = re.sub(r"\$\\times\s*1/255\$", "×1/255", tex)
    # drop layout-only commands pandoc cannot use
    for pat in (r"\\usepackage(\[[^\]]*\])?\{[^}]*\}", r"\\usetikzlibrary\{[^}]*\}", r"\\IEEEtriggeratref\{\d+\}",
                r"\\bibliographystyle\{[^}]*\}", r"\\setlength\{[^}]*\}\{[^}]*\}", r"\\renewcommand\{\\arraystretch\}\{[^}]*\}",
                r"\\addlinespace(\[[^\]]*\])?", r"\\footnotesize", r"\\centering"):
        tex = re.sub(pat, "", tex)
    # title block -> marked paragraphs
    title = re.search(r"\\title\{(.*?)\}\n", tex, re.S).group(1)
    names = re.search(r"\\IEEEauthorblockN\{(.*?)\}", tex, re.S).group(1)
    affil = re.search(r"\\IEEEauthorblockA\{(.*?)\}\n\}", tex, re.S).group(1).replace("\\\\", " \\newline ")
    tex = re.sub(r"\\title\{.*?\}\n", "", tex, flags=re.S)
    tex = re.sub(r"\\author\{.*?\n\}\n", "", tex, flags=re.S)
    tex = tex.replace("\\maketitle", f"{TITLE_MARK}{title}\n\n{AUTHOR_MARK}{names}\n\n{AFFIL_MARK}{affil}\n\n")
    # abstract / keywords in IEEE run-in style
    tex = re.sub(r"\\begin\{abstract\}\s*(.*?)\s*\\end\{abstract\}",
                 lambda m: "\\textbf{\\textit{Abstract}---" + m.group(1) + "}", tex, flags=re.S)
    tex = re.sub(r"\\begin\{IEEEkeywords\}\s*(.*?)\s*\\end\{IEEEkeywords\}",
                 lambda m: "\\textbf{\\textit{Index Terms}---" + m.group(1) + "}", tex, flags=re.S)
    # numbered headings: I., II. ... and A., B. ...
    out, sec, sub = [], 0, 0
    for line in tex.split("\n"):
        m = re.match(r"\\section\{(.*)\}\s*$", line)
        if m:
            sec += 1
            sub = 0
            line = f"\\section*{{{ROMAN[sec - 1]}. {m.group(1)}}}"
        m = re.match(r"\\subsection\{(.*)\}\s*$", line)
        if m:
            sub += 1
            line = f"\\subsection*{{{chr(64 + sub)}. {m.group(1)}}}"
        out.append(line)
    tex = "\n".join(out)
    # figures / tables: numbered captions, wide-float markers, PNG figures
    def float_env(m):
        env, body = m.group(1), m.group(2)
        wide = env.endswith("*")
        kind = "figure" if env.startswith("figure") else "table"
        lab = re.search(r"\\label\{([^}]+)\}", body)
        num = labels[lab.group(1)] if lab else "?"
        cap = (re.search(r"\\caption\{(.*?)\}\s*\\label\{[^}]*\}", body, re.S)
               or re.search(r"\\caption\{(.*)\}", body, re.S))
        cap_text = cap.group(1).strip()
        body = body.replace(cap.group(0), "")
        body = re.sub(r"\\label\{[^}]+\}", "", body)
        if "tikzpicture" in body:
            body = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}",
                          lambda _: "\\includegraphics{" + pipeline_png.replace("\\", "/") + "}", body, flags=re.S)
        body = re.sub(r"\\includegraphics(\[[^\]]*\])?\{figures/([^}]+)\.pdf\}",
                      lambda mm: "\\includegraphics{" + os.path.join(WORK, mm.group(2) + ".png").replace("\\", "/") + "}",
                      body)
        body = re.sub(r"\\includegraphics(\[[^\]]*\])?\{figures/([^}]+\.png)\}",
                      lambda mm: "\\includegraphics{" + os.path.join(HERE, "figures", mm.group(2)).replace("\\", "/") + "}",
                      body)
        if kind == "figure":
            block = f"{body.strip()}\n\n{FIGCAP_MARK}Fig. {num}. {cap_text}\n"
        else:
            block = f"{TABCAP_MARK}TABLE {num}\n\n{TABCAP_MARK}{cap_text}\n\n{body.strip()}\n"
        if wide:
            block = f"\n{WIDE_START}\n\n{block}\n\n{WIDE_END}\n"
        return "\n" + block + "\n"
    tex = re.sub(r"\\begin\{(figure\*?|table\*?)\}(?:\[[a-z]+\])?(.*?)\\end\{\1\}", float_env, tex, flags=re.S)
    # equations: keep numbers as a marked trailing paragraph
    counter = [0]

    def eq(m):
        counter[0] += 1
        return f"\\[{m.group(1).strip()}\\]\n\n{EQ_MARK}{counter[0]}@@\n"
    tex = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}", eq, tex, flags=re.S)
    # cross-references and citations
    tex = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels.get(m.group(1), "?"), tex)
    tex = re.sub(r"~?\\cite\{([^}]+)\}", lambda m: " " + format_cite(m.group(1), cites), tex)
    # bibliography
    tex = tex.replace("\\bibliography{refs}", "\\section*{References}\n\n" + bibliography_tex(cites))
    # small fixes
    tex = tex.replace("\\section*{Acknowledgment}", "\\section*{Acknowledgment}")
    tex = re.sub(r"%[^\n]*", "", tex.replace("\\%", "@@PCT@@")).replace("@@PCT@@", "\\%")
    return tex


# ----------------------------------------------------------------------------------------------
# 3. Reference document (IEEE-like styles)
# ----------------------------------------------------------------------------------------------
def reference_docx():
    ref = os.path.join(WORK, "ref.docx")
    data = pypandoc.convert_text("", "docx", format="md", outputfile=ref)  # noqa: F841  (creates default)
    doc = Document(ref)
    styles = doc.styles
    for st in styles:
        if st.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE):
            continue
        rPr = st.element.get_or_add_rPr()
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rPr.insert(0, fonts)
        for attr in list(fonts.attrib):  # drop theme fonts (e.g. Aptos/Calibri on headings)
            if "Theme" in attr or "theme" in attr:
                del fonts.attrib[attr]
        for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn(f"w:{slot}"), "Times New Roman")
        st.font.color.rgb = RGBColor(0, 0, 0)
    def para(name, size, bold=None, italic=None, align=None, before=0, after=0, first_line=None, small_caps=None):
        if name not in [s.name for s in styles]:
            return
        st = styles[name]
        st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold
        if italic is not None:
            st.font.italic = italic
        if small_caps is not None:
            st.font.small_caps = small_caps
        pf = st.paragraph_format
        if align is not None:
            pf.alignment = align
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.0
        if first_line is not None:
            pf.first_line_indent = first_line
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    para("Normal", 10, align=J)
    para("Body Text", 10, align=J, first_line=Inches(0.14))
    para("First Paragraph", 10, align=J, first_line=Inches(0.14))
    para("Compact", 10, align=J)
    para("Heading 1", 10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, small_caps=True)
    para("Heading 2", 10, bold=False, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3)
    para("Image Caption", 8, italic=False, align=J, before=3, after=8)
    para("Table Caption", 8, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2)
    para("Captioned Figure", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Figure", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    for h in ("Heading 1", "Heading 2"):
        if h in [s.name for s in styles]:
            styles[h].font.color.rgb = RGBColor(0, 0, 0)
    doc.save(ref)
    return ref


# ----------------------------------------------------------------------------------------------
# 4. Post-processing: layout, sections, captions, tables, references
# ----------------------------------------------------------------------------------------------
def set_cols(sectPr, n, space_in=0.25):
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_in * 1440)))


def page_setup(sectPr, section_type="continuous"):
    for tag in ("w:type", "w:pgSz", "w:pgMar"):
        el = sectPr.find(qn(tag))
        if el is not None:
            sectPr.remove(el)
    t = OxmlElement("w:type")
    t.set(qn("w:val"), section_type)
    sectPr.insert(0, t)
    sz = OxmlElement("w:pgSz")
    sz.set(qn("w:w"), str(int(8.5 * 1440)))
    sz.set(qn("w:h"), str(int(11 * 1440)))
    mar = OxmlElement("w:pgMar")
    for k, v in (("top", 0.75), ("bottom", 1.0), ("left", 0.625), ("right", 0.625), ("header", 0.3),
                 ("footer", 0.3), ("gutter", 0)):
        mar.set(qn(f"w:{k}"), str(int(v * 1440)))
    sectPr.append(sz)
    sectPr.append(mar)


def make_section_break(paragraph, cols, first=False):
    """Turn `paragraph` into the last paragraph of a section with `cols` columns."""
    pPr = paragraph._p.get_or_add_pPr()
    sectPr = OxmlElement("w:sectPr")
    page_setup(sectPr, "nextPage" if first else "continuous")
    set_cols(sectPr, cols)
    pPr.append(sectPr)
    for r in paragraph.runs:
        r.text = ""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def style_table(table, wide):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", 8), ("bottom", 8), ("insideH", 0), ("insideV", 0), ("left", 0), ("right", 0)):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)
    # header rule under the first row
    if len(table.rows):
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tb = OxmlElement("w:tcBorders")
            bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single")
            bot.set(qn("w:sz"), "4")
            bot.set(qn("w:color"), "000000")
            tb.append(bot)
            tcPr.append(tb)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(7.5 if wide else 7)
                    r.font.name = "Times New Roman"


def postprocess(path):
    doc = Document(path)
    body = doc.element.body
    paras = list(doc.paragraphs)

    # title block
    for p in paras:
        t = p.text
        if t.startswith(TITLE_MARK):
            p.text = t.replace(TITLE_MARK, "")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(10)
            for r in p.runs:
                r.font.size, r.font.name = Pt(24), "Times New Roman"
        elif t.startswith(AUTHOR_MARK):
            p.text = t.replace(AUTHOR_MARK, "")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            for r in p.runs:
                r.font.size = Pt(11)
        elif t.startswith(AFFIL_MARK):
            for r in p.runs:
                r.text = r.text.replace(AFFIL_MARK, "")
                r.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(14)
            title_end = p
    # one-column title section
    end_para = doc.add_paragraph()  # placeholder, replaced below
    end_para._p.getparent().remove(end_para._p)
    brk = title_end.insert_paragraph_before("")
    title_end._p.addnext(brk._p)
    make_section_break(brk, 1, first=True)

    # wide floats -> one-column sections
    for p in list(doc.paragraphs):
        if p.text.strip() == WIDE_START:
            make_section_break(p, 2)
        elif p.text.strip() == WIDE_END:
            make_section_break(p, 1)

    # equation numbers: append "(n)" to the preceding display-math paragraph
    for p in list(doc.paragraphs):
        m = re.fullmatch(r"@@EQNUM(\d+)@@", p.text.strip())
        if m:
            prev = p._p.getprevious()
            if prev is not None:
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = f"\t({m.group(1)})"
                r.append(t)
                prev.append(r)
                pPr = prev.find(qn("w:pPr"))
                if pPr is None:
                    pPr = OxmlElement("w:pPr")
                    prev.insert(0, pPr)
                tabs = OxmlElement("w:tabs")
                tab = OxmlElement("w:tab")
                tab.set(qn("w:val"), "right")
                tab.set(qn("w:pos"), str(int(3.43 * 1440)))
                tabs.append(tab)
                pPr.append(tabs)
            p._p.getparent().remove(p._p)

    # captions
    for p in doc.paragraphs:
        t = p.text
        if FIGCAP_MARK in t:
            for r in p.runs:
                r.text = r.text.replace(FIGCAP_MARK, "")
                r.font.size = Pt(8)
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if p.runs:
                m = re.match(r"(Fig\. \d+\.)(.*)", p.runs[0].text)
                if m:
                    p.runs[0].text = m.group(2)
                    lead = p.runs[0]._r
                    new = copy.deepcopy(lead)
                    lead.addprevious(new)
                    new.find(qn("w:t")).text = m.group(1)
        elif TABCAP_MARK in t:
            for r in p.runs:
                r.text = r.text.replace(TABCAP_MARK, "")
                r.font.size = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(6) if p.text.startswith("TABLE") else Pt(0)
            p.paragraph_format.space_after = Pt(0) if p.text.startswith("TABLE") else Pt(4)
            if not p.text.startswith("TABLE"):
                for r in p.runs:
                    r.font.small_caps = True
        elif t.startswith(REF_MARK):
            for r in p.runs:
                r.text = r.text.replace(REF_MARK, "")
                r.font.size = Pt(8)
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(0.3), Inches(-0.3)
            pf.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif t.startswith("Abstract") or t.startswith("Index Terms"):
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.size = Pt(9)

    # tables
    wide_ranges = []
    in_wide = False
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            sp = child.find(qn("w:pPr"))
            sect = sp.find(qn("w:sectPr")) if sp is not None else None
            if sect is not None:
                cols = sect.find(qn("w:cols"))
                in_wide = cols is not None and cols.get(qn("w:num")) == "2"
        elif child.tag == qn("w:tbl"):
            wide_ranges.append(in_wide)
    for table, wide in zip(doc.tables, wide_ranges):
        style_table(table, wide)

    # images: column width vs. full width
    in_wide = False
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            sp = child.find(qn("w:pPr"))
            sect = sp.find(qn("w:sectPr")) if sp is not None else None
            for ext in child.iter(qn("wp:extent")):
                cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                target = Inches(7.0) if in_wide else Inches(3.4)
                ext.set("cx", str(int(target)))
                ext.set("cy", str(int(target * cy / cx)))
                for a_ext in child.iter(qn("a:ext")):
                    a_ext.set("cx", str(int(target)))
                    a_ext.set("cy", str(int(target * cy / cx)))
            if sect is not None:
                cols = sect.find(qn("w:cols"))
                in_wide = cols is not None and cols.get(qn("w:num")) == "2"

    # final (body) section: two columns
    final = body.find(qn("w:sectPr"))
    page_setup(final, "continuous")
    set_cols(final, 2)
    doc.save(path)


def main():
    tex = read(os.path.join(HERE, "main.tex"))
    labels, cites = labels_and_cites()
    pipeline_png = build_tikz_png(tex)
    for f in os.listdir(os.path.join(HERE, "figures")):
        if f.endswith(".pdf"):
            render_pdf_png(os.path.join(HERE, "figures", f), os.path.join(WORK, f[:-4] + ".png"))
    pre = preprocess(tex, labels, cites, pipeline_png)
    src = os.path.join(WORK, "paper_docx.tex")
    with open(src, "w", encoding="utf-8") as f:
        f.write(pre)
    ref = reference_docx()
    pypandoc.convert_file(src, "docx", format="latex", outputfile=OUT,
                          extra_args=[f"--reference-doc={ref}", f"--resource-path={HERE}"])
    postprocess(OUT)
    print(f"wrote {OUT}  (intermediates in {WORK})")


if __name__ == "__main__":
    sys.exit(main())
